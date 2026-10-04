"""文档管理服务——上传、解析、入库、切分预览"""
import os
import uuid
import shutil
from datetime import datetime
from typing import List, Optional, BinaryIO

from sqlalchemy import select, func
from sqlalchemy.ext.asyncio import AsyncSession
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document as LangchainDocument

from app.models.document import Document, DocVersion, ChunkMeta, IngestTask, SecurityLevel
from app.schemas.document import DocumentCreate, DocumentUpdate
from app.config import get_settings
from app.services.vector_store import vector_store_service
from app.services.embedding_service import embedding_service

settings = get_settings()


class DocumentService:

    @staticmethod
    async def create_document(
        db: AsyncSession,
        data: DocumentCreate,
        user_dept: str = "通用部门",
    ) -> Document:
        doc = Document(
            kb_id=data.kb_id,
            title=data.title,
            author=data.author,
            dept=data.dept or user_dept,
            security_level=data.security_level,
            tags=data.tags,
        )
        db.add(doc)
        await db.flush()
        await db.refresh(doc)
        return doc

    @staticmethod
    async def get_document(db: AsyncSession, doc_id: str) -> Optional[Document]:
        result = await db.execute(select(Document).where(Document.id == doc_id))
        return result.scalar_one_or_none()

    @staticmethod
    async def list_documents(
        db: AsyncSession,
        kb_id: Optional[str] = None,
        status: str = "active",
        page: int = 1,
        page_size: int = 20,
    ) -> tuple[List[Document], int]:
        query = select(Document)
        count_q = select(func.count(Document.id))

        if kb_id:
            query = query.where(Document.kb_id == kb_id)
            count_q = count_q.where(Document.kb_id == kb_id)
        if status:
            query = query.where(Document.status == status)
            count_q = count_q.where(Document.status == status)

        total = (await db.execute(count_q)).scalar() or 0
        offset = (page - 1) * page_size
        query = query.order_by(Document.created_at.desc()).offset(offset).limit(page_size)
        result = await db.execute(query)
        items = list(result.scalars().all())

        return items, total

    @staticmethod
    async def update_document(db: AsyncSession, doc_id: str, data: DocumentUpdate) -> Optional[Document]:
        doc = await DocumentService.get_document(db, doc_id)
        if not doc:
            return None
        update_dict = data.model_dump(exclude_unset=True)
        for field, value in update_dict.items():
            setattr(doc, field, value)
        await db.flush()
        await db.refresh(doc)
        return doc

    @staticmethod
    async def delete_document(db: AsyncSession, doc_id: str) -> bool:
        doc = await DocumentService.get_document(db, doc_id)
        if not doc:
            return False
        await db.delete(doc)
        await db.flush()
        return True

    @staticmethod
    async def get_versions(db: AsyncSession, doc_id: str) -> List[DocVersion]:
        result = await db.execute(
            select(DocVersion)
            .where(DocVersion.doc_id == doc_id)
            .order_by(DocVersion.version.desc())
        )
        return list(result.scalars().all())

    @staticmethod
    async def upload_and_ingest(
        db: AsyncSession,
        doc_id: str,
        file: BinaryIO,
        filename: str,
        collection_name: str,
        kb_id: str,
    ) -> IngestTask:
        """
        上传文件并执行入库（同步版，MVP阶段简化）
        生产环境应异步到 Celery Worker
        """
        doc = await DocumentService.get_document(db, doc_id)
        if not doc:
            raise ValueError("文档不存在")

        # 保存文件
        upload_dir = os.path.abspath(settings.UPLOAD_DIR)
        os.makedirs(upload_dir, exist_ok=True)
        file_ext = os.path.splitext(filename)[1].lower()
        file_key = f"{kb_id}/{doc_id}/{uuid.uuid4().hex}{file_ext}"
        file_path = os.path.join(upload_dir, file_key)
        os.makedirs(os.path.dirname(file_path), exist_ok=True)

        with open(file_path, "wb") as f:
            shutil.copyfileobj(file, f)

        file_size = os.path.getsize(file_path)
        version = (doc.current_version or 0) + 1

        # 创建版本记录
        doc_version = DocVersion(
            doc_id=doc_id,
            version=version,
            oss_file_key=file_key,
            file_type=file_ext.lstrip("."),
            file_size=file_size,
            index_status="processing",
        )
        db.add(doc_version)
        await db.flush()

        # 创建入库任务
        task = IngestTask(
            doc_version_id=doc_version.id,
            status="processing",
        )
        db.add(task)
        await db.flush()

        try:
            # 解析文档文本
            text = DocumentService._parse_file(file_path, file_ext)

            # 切分
            kb = await db.execute(
                select(Document).where(Document.kb_id == kb_id)
            )
            # 使用默认切分参数
            splitter = RecursiveCharacterTextSplitter(
                chunk_size=800,
                chunk_overlap=150,
                separators=["\n\n", "\n", "。", "！", "？", ".", "!", "?", " ", ""],
            )
            chunks = splitter.split_text(text)

            # 构建 LangChain Document 列表
            langchain_docs = []
            chunk_metas = []
            for i, chunk_text in enumerate(chunks):
                chunk_id = f"{doc_id}:{version}:{i}"
                metadata = {
                    "chunk_id": chunk_id,
                    "kb_id": kb_id,
                    "doc_id": doc_id,
                    "doc_version": version,
                    "doc_title": doc.title,
                    "security_level": doc.security_level.value if isinstance(doc.security_level, SecurityLevel) else str(doc.security_level),
                    "department": doc.dept or "",
                    "page_no": i // 3 + 1,  # 粗略估算
                }
                langchain_docs.append(LangchainDocument(page_content=chunk_text, metadata=metadata))

                chunk_metas.append(ChunkMeta(
                    id=chunk_id,
                    doc_id=doc_id,
                    version=version,
                    chunk_index=i,
                    page_no=metadata["page_no"],
                    token_len=len(chunk_text),
                ))

            # 写入向量库
            vector_store_service.add_documents(collection_name, langchain_docs)

            # 写切片元数据到 MySQL
            for cm in chunk_metas:
                db.add(cm)

            # 更新状态
            doc_version.index_status = "success"
            doc_version.chunk_count = len(chunks)
            task.status = "success"
            task.progress = 100.0
            task.chunk_count = len(chunks)
            doc.current_version = version

            # 更新 KB 统计
            from app.models.kb import KnowledgeBase
            kb_result = await db.execute(select(KnowledgeBase).where(KnowledgeBase.id == kb_id))
            kb_obj = kb_result.scalar_one_or_none()
            if kb_obj:
                kb_obj.doc_count = (kb_obj.doc_count or 0) + 1
                kb_obj.chunk_count = (kb_obj.chunk_count or 0) + len(chunks)

        except Exception as e:
            doc_version.index_status = "failed"
            doc_version.error_msg = str(e)
            task.status = "failed"
            task.error = str(e)

        await db.flush()
        await db.refresh(task)
        return task

    @staticmethod
    def _parse_file(file_path: str, ext: str) -> str:
        """解析不同格式文档为纯文本"""
        if ext in (".txt", ".md", ".html", ".htm"):
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()

        elif ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            pages = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    pages.append(text)
            return "\n\n".join(pages)

        elif ext == ".docx":
            from docx import Document as DocxDocument
            doc = DocxDocument(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n".join(paragraphs)

        elif ext in (".xlsx", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            all_text = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                for row in ws.iter_rows(values_only=True):
                    row_text = " | ".join(str(cell) if cell is not None else "" for cell in row)
                    if row_text.strip():
                        all_text.append(row_text)
            return "\n".join(all_text)

        elif ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(file_path)
            slides = []
            for slide in prs.slides:
                texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        texts.append(shape.text)
                if texts:
                    slides.append("\n".join(texts))
            return "\n\n---\n\n".join(slides)

        else:
            raise ValueError(f"不支持的文件格式: {ext}")

    @staticmethod
    async def preview_chunks(
        db: AsyncSession,
        doc_id: str,
        chunk_size: int = 800,
        chunk_overlap: int = 150,
    ) -> dict:
        """切分预览——返回切片列表和预估 token 数"""
        doc = await DocumentService.get_document(db, doc_id)
        if not doc:
            raise ValueError("文档不存在")

        # 获取最新版本文件
        result = await db.execute(
            select(DocVersion)
            .where(DocVersion.doc_id == doc_id)
            .order_by(DocVersion.version.desc())
            .limit(1)
        )
        version = result.scalar_one_or_none()
        if not version:
            raise ValueError("文档尚无版本")

        file_path = os.path.join(settings.UPLOAD_DIR, version.oss_file_key)
        if not os.path.exists(file_path):
            raise ValueError("文件不存在")

        text = DocumentService._parse_file(file_path, f".{version.file_type}")
        splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
        )
        chunks = splitter.split_text(text)

        preview = []
        for i, chunk in enumerate(chunks[:20]):  # 最多预览 20 条
            preview.append({
                "index": i,
                "content": chunk[:300],
                "token_count": len(chunk),
                "page_no": i // 3 + 1,
            })

        return {
            "doc_id": doc_id,
            "total_chunks": len(chunks),
            "chunks": preview,
        }

    @staticmethod
    async def get_ingest_tasks(db: AsyncSession, doc_version_id: str) -> List[IngestTask]:
        result = await db.execute(
            select(IngestTask)
            .where(IngestTask.doc_version_id == doc_version_id)
            .order_by(IngestTask.created_at.desc())
        )
        return list(result.scalars().all())


# 单例
document_service = DocumentService()